"""Export scrape results to a formatted Excel file.

Handles both non-zoned retailers (Zone/Zip blank) and zoned retailers (Retailer M).
Output: OUTPUT/comp_scrape_YYYY-MM-DD.xlsx
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_DIR = Path(__file__).parent / "OUTPUT"

COLUMNS = [
    ("Competitor",       "competitor",          14),
    ("Zone",             "zone",                 6),
    ("Zip",              "zip_code",             8),
    ("Category",         "category",            16),
    ("Comp Product",      "comp_family",         22),
    ("Client Product",       "our_family",          22),
    ("Client Retail",       "our_retail",          12),
    ("List Price",       "comp_list_price",     12),
    ("Sale Price",       "comp_sale_price",     12),
    ("Active Price",     "comp_active_price",   13),
    ("Markdown %",       "markdown_pct",        12),
    ("vs Client Retail %",  "vs_our_retail_pct",   16),
    ("Price Changed",    "price_changed",       14),
    ("$ Change",         "price_change_delta",  12),
    ("Status",           "scrape_status",       10),
    ("URL",              "url",                 50),
    ("Notes",            "notes",               30),
]

# Fills
HEADER_FILL  = PatternFill("solid", fgColor="1F3864")
CHANGED_FILL = PatternFill("solid", fgColor="FFF2CC")
GREEN_FILL   = PatternFill("solid", fgColor="E2EFDA")
RED_FILL     = PatternFill("solid", fgColor="FCE4D6")
ERROR_FILL   = PatternFill("solid", fgColor="D9D9D9")
BLOCKED_FILL = PatternFill("solid", fgColor="FFE0CC")  # orange — bot-blocked
ALT_FILL     = PatternFill("solid", fgColor="F5F5F5")

HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
NORMAL_FONT = Font(size=10)

thin = Side(style="thin", color="CCCCCC")
THIN_BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def _write_sheet(ws, rows: list[dict]):
    ws.row_dimensions[1].height = 18
    for col_idx, (header, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill      = HEADER_FILL
        cell.font      = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        cell.border    = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for row_idx, row in enumerate(rows, start=2):
        is_changed  = row.get("price_changed", False)
        is_error    = row.get("scrape_status") == "error"
        is_no_price = row.get("scrape_status") == "no_price"
        is_blocked  = row.get("scrape_status") == "blocked"
        vs_pct      = row.get("vs_our_retail_pct")

        base_fill = (
            ERROR_FILL   if is_error   else
            BLOCKED_FILL if is_blocked else
            CHANGED_FILL if is_changed else
            ALT_FILL     if row_idx % 2 == 0 else None
        )

        for col_idx, (_, field, _) in enumerate(COLUMNS, start=1):
            value = row.get(field)

            if (is_no_price or is_blocked) and field in ("comp_list_price", "comp_sale_price", "comp_active_price"):
                value = "N/A"
            elif field == "scrape_status" and value == "no_price":
                value = "N/A"
            elif field == "price_changed":
                value = "YES" if value else ""
            elif field in ("markdown_pct", "vs_our_retail_pct") and value is not None:
                value = f"{value:+.1f}%"
            elif field == "price_change_delta" and value is not None:
                value = f"${value:+.2f}"

            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font      = NORMAL_FONT
            cell.border    = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(field == "notes"))

            if field == "vs_our_retail_pct" and vs_pct is not None and not is_error:
                cell.fill = GREEN_FILL if vs_pct > 0 else RED_FILL
            elif base_fill:
                cell.fill = base_fill

        ws.row_dimensions[row_idx].height = 15


def export_to_excel(results: list[dict], today: str, retailer_order: list[str] | None = None) -> Path:
    OUTPUT_DIR.mkdir(exist_ok=True)
    out_path = OUTPUT_DIR / f"comp_scrape_{today}.xlsx"

    # Determine retailer tab order: use provided order, falling back to appearance order in results
    present = {r["competitor"] for r in results}
    if retailer_order:
        ordered_retailers = [r for r in retailer_order if r in present]
        # Append any retailers not in the order list (e.g. new additions)
        ordered_retailers += [r for r in present if r not in ordered_retailers]
    else:
        # Preserve order of first appearance in results
        seen = []
        for r in results:
            if r["competitor"] not in seen:
                seen.append(r["competitor"])
        ordered_retailers = seen

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ── All Retailers summary sheet ───────────────────────────────────────
    ws_all = wb.create_sheet("All Retailers")
    _write_sheet(ws_all, results)

    # ── Per-retailer sheets ───────────────────────────────────────────────
    for retailer in ordered_retailers:
        retail_rows = [r for r in results if r["competitor"] == retailer]
        ws = wb.create_sheet(retailer[:31])
        _write_sheet(ws, retail_rows)

    # ── Per-zone sheets (only for zoned rows) ────────────────────────────
    zones = sorted({r["zone"] for r in results if r.get("zone") is not None})
    for zone in zones:
        zone_rows = sorted(
            [r for r in results if r.get("zone") == zone],
            key=lambda r: (r["competitor"], r["category"] or "", r["comp_family"] or ""),
        )
        ws = wb.create_sheet(f"Zone {zone}")
        _write_sheet(ws, zone_rows)

    # ── Price changes sheet (front) ───────────────────────────────────────
    changed = [r for r in results if r.get("price_changed")]
    if changed:
        ws_chg = wb.create_sheet("Price Changes")
        _write_sheet(ws_chg, changed)
        wb.move_sheet("Price Changes", offset=-(len(wb.sheetnames) - 1))

    wb.save(out_path)
    return out_path
