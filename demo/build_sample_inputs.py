"""
Build the sample input workbooks in sample_data/.

All product families and "client retail" prices are fictional. Replace the
sample files with your own (in RESOURCES/) to scrape real product pages.

Usage: python demo/build_sample_inputs.py
"""
from pathlib import Path

import openpyxl

OUT_DIR = Path(__file__).resolve().parent.parent / "sample_data"

NON_ZONED_HEADERS = ["Competitor", "Category", "Comp Product", "URL",
                     "Client Product", "Client Retail", "Notes/Insights"]
ZONED_HEADERS     = ["Competitor", "Zone", "Zip", "Category", "Comp Product", "URL",
                     "Client Product", "Client Retail", "Notes/Insights"]

# Retailer names are anonymized; URLs are placeholders
DOMAINS = {
    f"RETAILER {letter}": f"https://www.retailer-{letter.lower()}.example"
    for letter in "ABCDEFGHIJKL"
}

# (category, comp product, client product, client retail) — fictional
PRODUCTS = [
    ("Living Room",   "Comp Product 1", "Client Product 1", 1299.00),
    ("Living Room",   "Comp Product 2", "Client Product 2", 2199.00),
    ("Motion",        "Comp Product 3", "Client Product 3",  899.00),
    ("Dining Room",   "Comp Product 4", "Client Product 4", 1099.00),
    ("Adult Bedroom", "Comp Product 5", "Client Product 5",  749.00),
]

ZONES = [(1, "60601"), (2, "30301"), (3, "75201")]


def slug(text: str) -> str:
    return text.lower().replace(" ", "-")


def build_non_zoned():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Links"
    ws.append(NON_ZONED_HEADERS)
    for comp, domain in DOMAINS.items():
        for category, comp_family, our_family, our_retail in PRODUCTS:
            url = f"{domain}/REPLACE-WITH-PRODUCT-URL/{slug(comp_family)}"
            ws.append([comp, category, comp_family, url, our_family, our_retail, None])
    wb.save(OUT_DIR / "Non-Zoned Retailer Weekly Check Links.xlsx")


def build_zoned():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Links"
    ws.append(ZONED_HEADERS)
    for zone, zip_code in ZONES:
        for category, comp_family, our_family, our_retail in PRODUCTS[:3]:
            url = f"https://www.retailer-m.example/REPLACE-WITH-PRODUCT-URL/{slug(comp_family)}"
            ws.append(["RETAILER M", zone, zip_code, category, comp_family, url,
                       our_family, our_retail, None])
    wb.save(OUT_DIR / "Zoned Retailer Weekly Check Links.xlsx")


if __name__ == "__main__":
    OUT_DIR.mkdir(exist_ok=True)
    build_non_zoned()
    build_zoned()
    print(f"Sample inputs written to {OUT_DIR}")
