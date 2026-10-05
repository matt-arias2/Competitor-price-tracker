"""
Main scraper entry point — runs all retailers (non-zoned + zoned).

Usage:
    python run_scraper.py                         # all retailers
    python run_scraper.py --test                  # 2 rows per retailer/zone
    python run_scraper.py --retailer "Retailer E" "Retailer I"  # specific non-zoned retailers
    python run_scraper.py --retailer "Retailer M"               # zoned only
    python run_scraper.py --retailer "Retailer M" --zone 7      # specific zone
    python run_scraper.py --zone 1 3              # specific zones (zoned retailers)
"""
import argparse
from collections import defaultdict
from datetime import date
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright, Browser

import db
from excel_export import export_to_excel
from scrapers.retailer_a import RetailerAScraper
from scrapers.retailer_b import RetailerBScraper
from scrapers.retailer_c import RetailerCScraper
from scrapers.retailer_d import RetailerDScraper
from scrapers.retailer_e import RetailerEScraper
from scrapers.retailer_f import RetailerFScraper
from scrapers.retailer_g import RetailerGScraper
from scrapers.retailer_h import RetailerHScraper
from scrapers.retailer_i import RetailerIScraper
from scrapers.retailer_j import RetailerJScraper
from scrapers.retailer_k import RetailerKScraper
from scrapers.retailer_l import RetailerLScraper
from scrapers.retailer_m import RetailerMScraper

def _input_file(name: str) -> Path:
    """Prefer the private links file in RESOURCES/; fall back to the bundled sample."""
    root    = Path(__file__).parent
    private = root / "RESOURCES" / name
    return private if private.exists() else root / "sample_data" / name


NON_ZONED_FILE = _input_file("Non-Zoned Retailer Weekly Check Links.xlsx")
ZONED_FILE     = _input_file("Zoned Retailer Weekly Check Links.xlsx")


def get_file_retailer_order(filepath: Path) -> list[str]:
    """Return retailers in the order they first appear in the links file."""
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    seen = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        comp = row[0]
        if comp:
            comp = comp.strip().upper()
            if comp not in seen:
                seen.append(comp)
    return seen


def get_file_url_order(filepath: Path) -> dict[str, int]:
    """Return {url: file_row_index} for every URL in the links file."""
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    order = {}
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        url = row[3]  # column D
        if url:
            order[url.strip()] = idx
    return order

UA_CHROMIUM = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
UA_FIREFOX  = "Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:146.0) Gecko/20100101 Firefox/146.0"

SCRAPERS = {
    "RETAILER A": RetailerAScraper(),
    "RETAILER B": RetailerBScraper(),
    "RETAILER C": RetailerCScraper(),
    "RETAILER D": RetailerDScraper(),
    "RETAILER E": RetailerEScraper(),
    "RETAILER F": RetailerFScraper(),
    "RETAILER G": RetailerGScraper(),
    "RETAILER H": RetailerHScraper(),
    "RETAILER I": RetailerIScraper(),
    "RETAILER J": RetailerJScraper(),
    "RETAILER K": RetailerKScraper(),
    "RETAILER L": RetailerLScraper(),
}

ZONED_SCRAPERS = {
    "RETAILER M": RetailerMScraper(),
}

PAGE_WAIT_MS = 4000


# ── Input loaders ─────────────────────────────────────────────────────────────

def load_non_zoned(
    filter_retailers: list[str] | None = None,
    test_mode: bool = False,
) -> list[dict]:
    wb = openpyxl.load_workbook(NON_ZONED_FILE)
    ws = wb.active
    rows = []
    counts: dict[str, int] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        comp, category, comp_family, url, our_family, our_retail, notes = row
        if not url or not comp:
            continue
        comp = comp.strip().upper()
        if filter_retailers and comp not in filter_retailers:
            continue
        if test_mode:
            counts[comp] = counts.get(comp, 0) + 1
            if counts[comp] > 2:
                continue
        rows.append({
            "competitor":  comp,
            "zone":        None,
            "zip_code":    None,
            "category":    category,
            "comp_family": comp_family.strip() if comp_family else "",
            "url":         url.strip(),
            "our_family":  our_family.strip() if our_family else "",
            "our_retail":  float(our_retail) if our_retail else None,
            "notes":       notes,
        })
    return rows


def load_zoned(
    filter_retailers: list[str] | None = None,
    filter_zones: list[int] | None = None,
    test_mode: bool = False,
) -> list[dict]:
    wb = openpyxl.load_workbook(ZONED_FILE)
    ws = wb.active
    rows = []
    counts: dict[tuple, int] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        comp, zone, zip_code, category, comp_family, url, our_family, our_retail, notes = row
        if not url or not comp or zone is None:
            continue
        comp = comp.strip().upper()
        zone = int(zone)
        if comp not in ZONED_SCRAPERS:
            continue  # skip Retailer N and any other unimplemented retailers
        if filter_retailers and comp not in filter_retailers:
            continue
        if filter_zones and zone not in filter_zones:
            continue
        if test_mode:
            key = (comp, zone)
            counts[key] = counts.get(key, 0) + 1
            if counts[key] > 2:
                continue
        rows.append({
            "competitor":  comp,
            "zone":        zone,
            "zip_code":    str(zip_code).strip(),
            "category":    category,
            "comp_family": comp_family.strip() if comp_family else "",
            "url":         url.strip(),
            "our_family":  our_family.strip() if our_family else "",
            "our_retail":  float(our_retail) if our_retail else None,
            "notes":       notes,
        })
    return rows


# ── Result builder ────────────────────────────────────────────────────────────

def build_result(row: dict, prices: dict, today: str, prev_active) -> dict:
    list_p = prices.get("comp_list_price")
    sale_p = prices.get("comp_sale_price")
    active = sale_p if sale_p else list_p

    markdown_pct = None
    if list_p and sale_p and list_p > 0:
        markdown_pct = round((list_p - sale_p) / list_p * 100, 1)

    vs_retail = None
    our_retail = row.get("our_retail")
    if active and our_retail and our_retail > 0:
        vs_retail = round((active - our_retail) / our_retail * 100, 1)

    price_changed = (prev_active is not None) and (active is not None) and (round(prev_active, 2) != round(active, 2))
    price_delta   = round(active - prev_active, 2) if (price_changed and active and prev_active) else None
    no_price      = (list_p is None and sale_p is None)

    return {
        **row,
        "scraped_date":       today,
        "comp_list_price":    list_p,
        "comp_sale_price":    sale_p,
        "comp_active_price":  active,
        "markdown_pct":       markdown_pct,
        "vs_our_retail_pct":  vs_retail,
        "price_changed":      price_changed,
        "price_change_delta": price_delta,
        "scrape_status":      "no_price" if no_price else "ok",
        "error_msg":          None,
    }


# ── Scrapers ──────────────────────────────────────────────────────────────────

def scrape_non_zoned(rows: list[dict], today: str) -> list[dict]:
    prev_prices = db.get_last_run_prices(exclude_date=today)
    results     = []

    # Tag each row with its original file position so we can restore order after batching
    for i, row in enumerate(rows):
        row["_idx"] = i

    http_rows     = [r for r in rows if SCRAPERS.get(r["competitor"]) and
                     getattr(SCRAPERS[r["competitor"]], "use_http", False)]
    chromium_rows = [r for r in rows if SCRAPERS.get(r["competitor"]) and
                     not getattr(SCRAPERS[r["competitor"]], "use_http", False) and
                     getattr(SCRAPERS[r["competitor"]], "browser_type", None) != "firefox"]
    firefox_rows  = [r for r in rows if SCRAPERS.get(r["competitor"]) and
                     not getattr(SCRAPERS[r["competitor"]], "use_http", False) and
                     getattr(SCRAPERS[r["competitor"]], "browser_type", None) == "firefox"]

    def run_http_batch(batch: list[dict]):
        for row in batch:
            scraper = SCRAPERS.get(row["competitor"])
            if not scraper:
                continue
            print(f"  Scraping {row['competitor']:15s} | {row['comp_family']:25s} | {row['url'][:60]}")
            try:
                prices = scraper.fetch_and_extract(row["url"])
                prev   = prev_prices.get(row["url"])
                result = build_result(row, prices, today, prev)
                results.append(result)
                print(f"    -> list={prices.get('comp_list_price')}  sale={prices.get('comp_sale_price')}  [{result['scrape_status']}]")
            except Exception as e:
                print(f"    -> ERROR: {e}")
                err = build_result(row, {}, today, None)
                err["scrape_status"] = "error"
                err["error_msg"]     = str(e)[:200]
                results.append(err)

    def try_scrape(br: Browser, ua: str, scraper, row: dict):
        """Single attempt: load the page and extract prices. Raises on error."""
        page = br.new_page(user_agent=ua)
        try:
            page.goto(row["url"], wait_until="domcontentloaded", timeout=30000)
            page.wait_for_timeout(PAGE_WAIT_MS)
            if "Access Denied" in (page.title() or ""):
                return {"_blocked": True}
            return scraper.extract(page, row["url"])
        finally:
            page.close()

    def run_batch(batch: list[dict], browser_type: str):
        ua = UA_FIREFOX if browser_type == "firefox" else UA_CHROMIUM
        with sync_playwright() as p:
            br: Browser = getattr(p, browser_type).launch(headless=True)
            for row in batch:
                scraper = SCRAPERS.get(row["competitor"])
                if not scraper:
                    continue
                print(f"  Scraping {row['competitor']:15s} | {row['comp_family']:25s} | {row['url'][:60]}")

                prices = None
                exc = None
                # One retry — covers transient timeouts and pages that render
                # price widgets a beat too late for the fixed post-load wait.
                for attempt in range(2):
                    try:
                        prices = try_scrape(br, ua, scraper, row)
                        exc = None
                    except Exception as e:
                        exc = e
                        prices = None

                    failed = (exc is not None or prices is None or prices.get("_blocked") or
                              (not prices.get("comp_list_price") and not prices.get("comp_sale_price")))
                    if not failed or attempt == 1:
                        break

                if exc is not None:
                    print(f"    -> ERROR: {exc}")
                    err = build_result(row, {}, today, None)
                    err["scrape_status"] = "error"
                    err["error_msg"]     = str(exc)[:200]
                    results.append(err)
                    continue

                if prices.get("_blocked"):
                    print(f"    -> BLOCKED (bot detection / Akamai)")
                    blocked = build_result(row, {}, today, None)
                    blocked["scrape_status"] = "blocked"
                    blocked["error_msg"]     = "Access Denied — bot detection active"
                    results.append(blocked)
                    continue

                prev   = prev_prices.get(row["url"])
                result = build_result(row, prices, today, prev)
                results.append(result)
                print(f"    -> list={prices.get('comp_list_price')}  sale={prices.get('comp_sale_price')}  [{result['scrape_status']}]")
            br.close()

    print(f"\nHTTP batch ({len(http_rows)} URLs)")
    run_http_batch(http_rows)
    print(f"\nChromium batch ({len(chromium_rows)} URLs)")
    run_batch(chromium_rows, "chromium")
    print(f"\nFirefox batch ({len(firefox_rows)} URLs)")
    run_batch(firefox_rows, "firefox")

    # Restore file order, then strip the temporary index key
    results.sort(key=lambda r: r.get("_idx", 0))
    for r in results:
        r.pop("_idx", None)
    return results


def scrape_zoned(rows: list[dict], today: str) -> list[dict]:
    prev_prices = db.get_last_run_zoned_prices(exclude_date=today)
    results     = []

    groups: dict[tuple, list[dict]] = defaultdict(list)
    for row in rows:
        groups[(row["competitor"], row["zone"], row["zip_code"])].append(row)

    with sync_playwright() as p:
        browser: Browser = p.chromium.launch(headless=True)

        for (competitor, zone, zip_code), zone_rows in sorted(groups.items()):
            scraper = ZONED_SCRAPERS.get(competitor)
            if not scraper:
                continue
            print(f"\n--- {competitor}  Zone {zone}  (zip {zip_code})  ---  {len(zone_rows)} URLs ---")

            context = browser.new_context(user_agent=UA_CHROMIUM)
            print(f"  Setting zip {zip_code}...")
            try:
                scraper.setup_context(context, zip_code)
                print(f"  Zip set OK")
            except Exception as e:
                print(f"  WARNING: setup_context failed: {e}")

            for row in zone_rows:
                print(f"  Scraping {competitor:8s} Z{zone} | {row['comp_family']:25s} | {row['url'][:60]}")
                page = context.new_page()
                try:
                    page.goto(row["url"], wait_until="domcontentloaded", timeout=30000)
                    page.wait_for_timeout(PAGE_WAIT_MS)
                    prices = scraper.extract(page, row["url"])
                    prev   = prev_prices.get((row["url"], row["zone"], row["zip_code"]))
                    result = build_result(row, prices, today, prev)
                    results.append(result)
                    print(f"    -> list={prices.get('comp_list_price')}  sale={prices.get('comp_sale_price')}  [{result['scrape_status']}]")
                except Exception as e:
                    print(f"    -> ERROR: {e}")
                    err = build_result(row, {}, today, None)
                    err["scrape_status"] = "error"
                    err["error_msg"]     = str(e)[:200]
                    results.append(err)
                finally:
                    page.close()
            context.close()

        browser.close()
    return results


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--test",     action="store_true", help="2 rows per retailer/zone")
    parser.add_argument("--retailer", nargs="+",           help="Limit to specific retailer(s)")
    parser.add_argument("--zone",     nargs="+", type=int, help="Limit to specific zone number(s) (zoned retailers only)")
    args = parser.parse_args()

    today = date.today().isoformat()
    print(f"=== Comp Scraper --- {today} ===")

    db.init_db()
    db.init_zoned_db()

    filter_r         = [r.upper() for r in args.retailer] if args.retailer else None
    non_zoned_names  = set(SCRAPERS.keys())
    zoned_names      = set(ZONED_SCRAPERS.keys())
    run_non_zoned    = not filter_r or bool(non_zoned_names & set(filter_r))
    run_zoned        = not filter_r or bool(zoned_names & set(filter_r))

    all_results: list[dict] = []

    # ── Non-zoned retailers ──────────────────────────────────────────────
    if run_non_zoned:
        nz_filter = [r for r in filter_r if r in non_zoned_names] if filter_r else None
        nz_rows   = load_non_zoned(filter_retailers=nz_filter, test_mode=args.test)
        print(f"\nNon-zoned: {len(nz_rows)} URLs to scrape")
        nz_results = scrape_non_zoned(nz_rows, today)

        db_rows = [{k: v for k, v in r.items()
                    if k not in ("price_changed", "price_change_delta", "zone", "zip_code")}
                   for r in nz_results]
        db.insert_results(db_rows)
        print(f"\nSaved {len(db_rows)} non-zoned rows to database")
        all_results.extend(nz_results)

    # ── Zoned retailers (Retailer M) ────────────────────────────────────────────
    if run_zoned:
        z_filter  = [r for r in filter_r if r in zoned_names] if filter_r else None
        z_rows    = load_zoned(filter_retailers=z_filter, filter_zones=args.zone, test_mode=args.test)
        if z_rows:
            print(f"\nZoned: {len(z_rows)} URLs to scrape")
            z_results = scrape_zoned(z_rows, today)

            db_rows_z = [{k: v for k, v in r.items()
                          if k not in ("price_changed", "price_change_delta")}
                         for r in z_results]
            db.insert_zoned_results(db_rows_z)
            print(f"\nSaved {len(db_rows_z)} zoned rows to database")
            all_results.extend(z_results)

    # ── Export combined Excel ────────────────────────────────────────────
    if all_results:
        retailer_order = get_file_retailer_order(NON_ZONED_FILE)
        out_path = export_to_excel(all_results, today, retailer_order=retailer_order)
        print(f"\nExcel exported: {out_path}")

    # ── Summary ──────────────────────────────────────────────────────────
    changed   = [r for r in all_results if r.get("price_changed")]
    errors    = [r for r in all_results if r.get("scrape_status") == "error"]
    no_prices = [r for r in all_results if r.get("scrape_status") == "no_price"]
    blocked   = [r for r in all_results if r.get("scrape_status") == "blocked"]
    print(f"\n-- Summary --")
    print(f"  Total scraped : {len(all_results)}")
    print(f"  Price changes : {len(changed)}")
    print(f"  N/A (OOS)     : {len(no_prices)}")
    print(f"  Blocked       : {len(blocked)}")
    print(f"  Errors        : {len(errors)}")
    if blocked:
        retailers = sorted({r['competitor'] for r in blocked})
        print(f"  Blocked retailers: {', '.join(retailers)}")

    if changed:
        print("\n  Price Changes:")
        for r in changed:
            zone_tag = f" Z{r['zone']}" if r.get("zone") else ""
            delta    = f"+{r['price_change_delta']}" if r['price_change_delta'] > 0 else str(r['price_change_delta'])
            print(f"    {r['competitor']:15s}{zone_tag:4s} {r['comp_family']:25s} {delta:>8s}  {r['url'][:50]}")

    if errors:
        print("\n  Errors:")
        for r in errors:
            zone_tag = f" Z{r['zone']}" if r.get("zone") else ""
            print(f"    {r['competitor']:15s}{zone_tag:4s} {r['comp_family']:25s}  {r['error_msg'][:60]}")


if __name__ == "__main__":
    main()
